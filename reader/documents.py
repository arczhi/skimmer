"""Document import: extract plain text from uploaded files.

Supported:
  .txt / .md / .markdown      plain decode (utf-8 with fallbacks)
  .pdf                        pypdf (pure python)
  .docx                       python-docx (paragraphs + tables)
  .doc                        best effort: antiword binary if present,
                              otherwise a clear error message

Everything runs locally; uploaded bytes never leave the machine.
"""

from __future__ import annotations

import io
import re
import shutil
import subprocess
import tempfile
from pathlib import Path


class DocumentError(Exception):
    pass


SUPPORTED = (".txt", ".md", ".markdown", ".pdf", ".docx", ".doc")


def _clean(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _decode_plain(data: bytes) -> str:
    for enc in ("utf-8", "utf-8-sig", "gb18030", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def extract_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as e:
        raise DocumentError("解析 PDF 需要 pypdf：pip install pypdf") from e
    reader = PdfReader(io.BytesIO(data))
    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:  # noqa: BLE001 - a broken page should not kill the import
            pages.append("")
    text = "\n\n".join(p for p in pages if p.strip())
    if not text.strip():
        raise DocumentError("PDF 未提取到文本（可能是扫描件/纯图片，需要 OCR）")
    return text


def extract_docx(data: bytes) -> str:
    try:
        import docx  # python-docx
    except ImportError as e:
        raise DocumentError("解析 DOCX 需要 python-docx：pip install python-docx") from e
    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    text = "\n".join(parts)
    if not text.strip():
        raise DocumentError("DOCX 未提取到文本")
    return text


def extract_doc(data: bytes) -> str:
    """Legacy .doc: use antiword when installed, otherwise fail with guidance."""
    antiword = shutil.which("antiword")
    if not antiword:
        raise DocumentError("旧版 .doc 需要 antiword：apt install antiword / brew install antiword；或先另存为 .docx")
    with tempfile.NamedTemporaryFile(suffix=".doc", delete=True) as f:
        f.write(data)
        f.flush()
        try:
            out = subprocess.run(
                [antiword, f.name], capture_output=True, timeout=60, check=False
            )
        except subprocess.TimeoutExpired as e:
            raise DocumentError("antiword 超时") from e
    if out.returncode != 0 or not out.stdout.strip():
        raise DocumentError("antiword 解析失败（文件可能损坏或受密码保护）")
    return _decode_plain(out.stdout)


def extract_text(filename: str, data: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED:
        raise DocumentError(f"暂不支持 {ext or '该'} 格式；支持：{', '.join(SUPPORTED)}")
    if ext == ".pdf":
        text = extract_pdf(data)
    elif ext == ".docx":
        text = extract_docx(data)
    elif ext == ".doc":
        text = extract_doc(data)
    else:
        text = _decode_plain(data)
    text = _clean(text)
    if not text:
        raise DocumentError("文件解析结果为空")
    return text
